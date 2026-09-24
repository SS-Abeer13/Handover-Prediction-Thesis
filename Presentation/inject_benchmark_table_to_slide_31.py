import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PPTX_PATH = r"D:\Handover Thesis\Presentation\Handover-Thesis-Defence.pptx"
prs = Presentation(PPTX_PATH)

slide_idx = 30 # 31st slide (0-indexed 30)
slide = prs.slides[slide_idx]

# Identify and remove old body shapes except background picture
bg_shape = None
shapes_to_remove = []
for i, s in enumerate(slide.shapes):
    if i == 0 and s.shape_type == 13: # background picture
        bg_shape = s
    else:
        shapes_to_remove.append(s)

print(f"Removing {len(shapes_to_remove)} old shapes from slide 31...")
for s in shapes_to_remove:
    sp = s._element
    sp.getparent().remove(sp)

# Styling Constants
MAROON = RGBColor(0x89, 0x13, 0x13)
NAVY = RGBColor(0x1F, 0x4E, 0x79)
INK = RGBColor(0x26, 0x26, 0x26)
MUTED = RGBColor(0x77, 0x77, 0x77)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HIGHLIGHT_BG = RGBColor(0xFD, 0xED, 0xEC)
HEADER_BG = RGBColor(0x1F, 0x4E, 0x79)
ALT_BG = RGBColor(0xF8, 0xF9, 0xFA)
BORDER_GRAY = RGBColor(0xCC, 0xCC, 0xCC)
FONT = "Arial"
L = 0.92

# 1. Breadcrumb
tb_bc = slide.shapes.add_textbox(Inches(L), Inches(0.85), Inches(7.0), Inches(0.32))
tf_bc = tb_bc.text_frame
tf_bc.word_wrap = True
p_bc = tf_bc.paragraphs[0]
p_bc.text = "LITERATURE BENCHMARK"
p_bc.font.size = Pt(11)
p_bc.font.bold = True
p_bc.font.name = FONT
p_bc.font.color.rgb = MUTED

# 2. Action Title
tb_title = slide.shapes.add_textbox(Inches(L), Inches(1.18), Inches(9.0), Inches(0.95))
tf_title = tb_title.text_frame
tf_title.word_wrap = True
p_title = tf_title.paragraphs[0]
p_title.text = "Benchmarked across 22 studies: our framework is the first to combine drive-level holdouts, multi-horizon lead time, and calibrated risk"
p_title.font.size = Pt(20)
p_title.font.bold = True
p_title.font.name = FONT
p_title.font.color.rgb = MAROON

# 3. Main Benchmark Table (Left Side, 8.35 inches wide)
tbl_left = Inches(L)
tbl_top = Inches(2.25)
tbl_w = Inches(8.35)
tbl_h = Inches(4.15)

rows = 6
cols = 6
col_widths = [Inches(1.55), Inches(1.30), Inches(1.35), Inches(1.65), Inches(1.35), Inches(1.15)]

table_shape = slide.shapes.add_table(rows, cols, tbl_left, tbl_top, tbl_w, tbl_h)
tbl = table_shape.table
for idx, w in enumerate(col_widths):
    tbl.columns[idx].width = w

headers = ["Study & Venue", "Data Source", "Split Protocol", "Primary Metric", "Calibration & Risk", "Lead Time"]
table_data = [
    [
        "Boutiba et al. (2021)\nIEEE GLOBECOM",
        "5G Testbed (OAI)\nSimulated fading",
        "Unspecified\n(Leakage risk)",
        "98.03% Accuracy\n(on rare RLF)",
        "None\n(Uncalibrated)",
        "None\n(Fixed 5 steps)"
    ],
    [
        "Dzaferagic (2024)\nIEEE TNSM",
        "Real O-RAN\n4,350 HOs (10.9%)",
        "Not Stated\n(Random row)",
        "Precision / Recall\n(at single point)",
        "None\n(Uncalibrated)",
        "None\n(Fixed horizon)"
    ],
    [
        "Shafi et al. (2025)\nMendeley / IUT",
        "Real XCAL (Dhaka)\n2,154 rows, 367 HOs",
        "Day 1 → Day 2\n(Single test run)",
        "HO Count reduced\n(AUROC 0.49 imminence)",
        "None\n(Heuristic reward)",
        "None\n(Reactive only)"
    ],
    [
        "Amirova et al. (2026)\nFuture Internet",
        "Real LTE (Astana)\n27k rows, 232 HOs",
        "Device hold-out\n(Non-random)",
        "PR curve, AUROC 0.82\n(No lift reported)",
        "Bootstrap AUC CI\n(No ECE / Brier)",
        "None\n(Single horizon)"
    ],
    [
        "This Work (2026)\nIUT EEE / CSE",
        "Real LTE (Dhaka)\n4 campaigns, 957 HOs",
        "Grouped whole-drive\n+ locked route holdout",
        "AUPRC 0.19–0.49 (3.4× lift)\nAUROC 0.816–0.834",
        "ECE 0.037, Brier 0.054\nConformal risk (≥89%)",
        "1.0 s to 5.0 s\n(Mean 3.45 s @ 5s)"
    ]
]

# Style Header Row
for c_idx, title in enumerate(headers):
    cell = tbl.cell(0, c_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = HEADER_BG
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_top = Inches(0.04)
    cell.margin_bottom = Inches(0.04)
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.CENTER
    p.font.bold = True
    p.font.size = Pt(10.5)
    p.font.name = FONT
    p.font.color.rgb = WHITE

# Style Data Rows
for r_idx, row_vals in enumerate(table_data):
    is_ours = (r_idx == len(table_data) - 1)
    bg = HIGHLIGHT_BG if is_ours else (ALT_BG if r_idx % 2 == 1 else WHITE)
    
    for c_idx, val in enumerate(row_vals):
        cell = tbl.cell(r_idx + 1, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_top = Inches(0.04)
        cell.margin_bottom = Inches(0.04)
        cell.margin_left = Inches(0.06)
        cell.margin_right = Inches(0.06)
        tf = cell.text_frame
        tf.word_wrap = True
        
        lines = val.split("\n")
        p = tf.paragraphs[0]
        p.text = lines[0]
        p.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER
        p.font.name = FONT
        p.font.size = Pt(9.5)
        p.font.bold = is_ours
        p.font.color.rgb = MAROON if is_ours else INK
        
        if len(lines) > 1:
            p2 = tf.add_paragraph()
            p2.text = lines[1]
            p2.alignment = PP_ALIGN.LEFT if c_idx == 0 else PP_ALIGN.CENTER
            p2.font.name = FONT
            p2.font.size = Pt(8.5)
            p2.font.italic = not is_ours
            p2.font.color.rgb = MAROON if is_ours else MUTED

# 4. Right Side: Interpretive Analysis Card
card_left = Inches(9.45)
card_top = Inches(2.25)
card_w = Inches(2.95)
card_h = Inches(4.15)

tb_card = slide.shapes.add_textbox(card_left, card_top, card_w, card_h)
tf_card = tb_card.text_frame
tf_card.word_wrap = True

p_ct = tf_card.paragraphs[0]
p_ct.text = "METHODOLOGICAL GAPS IN LITERATURE"
p_ct.font.bold = True
p_ct.font.size = Pt(11)
p_ct.font.name = FONT
p_ct.font.color.rgb = MAROON
p_ct.space_after = Pt(8)

bullets_data = [
    ("Why headline accuracies fail to travel: ", "7 of 22 audited papers report 94–99% accuracy on rare events. A naive model predicting 'no handover' achieves 93.3% accuracy here while detecting zero events."),
    ("Severe temporal data leakage: ", "10 of 22 papers omit split details. Random-row splitting inflates GRU AUPRC by +92% due to adjacent channel correlation. Grouped-drive holdout is mandatory."),
    ("First calibrated risk guarantees: ", "0 of 22 papers report calibration curves, ECE, Brier score, or lead time. Our framework guarantees bounded risk (coverage ≥ 89.1%) with 3.45 s advance warning.")
]

for lead, body in bullets_data:
    p = tf_card.add_paragraph()
    p.space_after = Pt(8)
    r1 = p.add_run()
    r1.text = lead
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.name = FONT
    r1.font.color.rgb = MAROON
    
    r2 = p.add_run()
    r2.text = body
    r2.font.size = Pt(9.5)
    r2.font.name = FONT
    r2.font.color.rgb = INK

# 5. Footnote / In-Slide Citation
tb_cite = slide.shapes.add_textbox(Inches(L), Inches(6.58), Inches(11.5), Inches(0.35))
tf_cite = tb_cite.text_frame
tf_cite.word_wrap = True
p_cite = tf_cite.paragraphs[0]
p_cite.text = "Systematic protocol audit of 22 published handover prediction models across 108 screened references (pass 3). Full 22-row audit matrix in thesis Chapter 2."
p_cite.font.size = Pt(10)
p_cite.font.name = FONT
p_cite.font.color.rgb = MUTED

prs.save(PPTX_PATH)
print("Successfully injected benchmark table and gaps card into Slide 31!")
