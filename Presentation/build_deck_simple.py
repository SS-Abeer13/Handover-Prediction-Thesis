"""Build the short, format-mandated thesis presentation on the IUT template.

Eleven sections, as required:
  1 Title page               7 Knowledge profile, complex problems and activities
  2 Outline                  8 Impact on health, society, environment, sustainability
  3 Background, motivation   9 Results and discussions
  4 Research methodology    10 Conclusion
  5 Proposed system         11 Future works
  6 Theory and experiments

Kept deliberately simple: plain language, one idea per slide, and nothing that
cannot be defended from the measurements themselves.
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deckkit import *                      # noqa: F401,F403  (template + slide helpers)
from deckkit import (new_deck, add_slide, head, textbox, para, bullets, cite,
                     callout, figure, fig_slide, flow_slide, bullet_slide,
                     table_slide, takeaway_bar, L, BODY_TOP, MAROON, INK, MUTED)
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

OUT = "/mnt/user-data/outputs/Handover-Thesis-Presentation.pptx"
prs = new_deck()

# ---------------------------------------------------------------- 1. TITLE
s = add_slide(prs, BG_TITLE)
tf = textbox(s, 0.80, 1.05, 8.2, 1.5)
para(tf, "Predicting the Next LTE Handover\nfrom Drive-Test Signalling",
     31, MAROON, bold=True, first=True, space_after=0)
tf = textbox(s, 0.80, 2.52, 8.6, 0.4)
para(tf, "How early can a handover be seen, and how far can the answer be trusted?",
     15, INK, first=True, italic=True, space_after=0)
tf = textbox(s, 0.80, 3.02, 8.6, 1.4)
para(tf, "Abeer Saadman", 19, INK, bold=True, first=True, space_after=2)
para(tf, "Student ID [ID]", 13, MUTED, space_after=8)
para(tf, "Supervisor: [Supervisor Name, Title]", 15, INK, space_after=6)
para(tf, "Department of Electrical and Electronic Engineering", 13, MUTED, space_after=1)
para(tf, "Islamic University of Technology, Gazipur, Bangladesh", 13, MUTED, space_after=1)
tf = textbox(s, 0.80, 4.20, 8.6, 0.4)
para(tf, "B.Sc. Thesis Defence   |   [Date]", 14, MAROON, bold=True, first=True, space_after=0)

# ---------------------------------------------------------------- 2. OUTLINE
s = add_slide(prs)
head(s, "Outline", "Outline of the presentation")
left = ["1.  Title page",
        "2.  Outline of the presentation",
        "3.  Background, motivation and objectives",
        "4.  Research methodology",
        "5.  Proposed system",
        "6.  Theory, modelling and experiments"]
right = ["7.  Knowledge profile and complex problems",
         "8.  Impact on society and sustainability",
         "9.  Results and discussions",
         "10.  Conclusion",
         "11.  Future works"]
tf = textbox(s, L + 0.25, BODY_TOP + 0.10, 5.4, 3.5)
for i, t in enumerate(left):
    para(tf, t, 19, INK, bold=True, first=(i == 0), space_after=16)
tf = textbox(s, 6.85, BODY_TOP + 0.10, 5.5, 3.5)
for i, t in enumerate(right):
    para(tf, t, 19, INK, bold=True, first=(i == 0), space_after=16)
takeaway_bar(s, "Four things to take away: the problem, the system, the evidence, and the boundary of the claim.")

# ============================ 3. BACKGROUND, MOTIVATION, OBJECTIVES =========
fig_slide(prs, "3 · Background",
    "LTE hands over using a rule that can only react",
    "fig01_a3_event",
    [("The phone measures ", "its own cell and its neighbours, once per second."),
     ("Event A3 fires ", "when a neighbour becomes better by a fixed offset and stays better for a fixed time."),
     ("Only then ", "is a report sent, and only then does the network decide."),
     ("On this network ", "the offset is +1 dB and the time-to-trigger is 320 ms.")],
    takeaway="By definition the rule cannot act until the radio has already gone bad.",
    cite_text="Entering condition and time-to-trigger per 3GPP TS 36.331. Parameters read from the operator's own configuration messages.")

fig_slide(prs, "3 · Motivation",
    "The cost of reacting late is measurable on our own data",
    "fig27_necessity",
    [("938 handovers in 2.9 hours ", "- about one every eleven seconds."),
     ("A quarter go straight back ", "to the cell just left. Those are wasted."),
     ("The link dropped 341 times ", "and had to be re-established."),
     ("And nobody has built the predictor: ", "of 22 comparable published models, none is tested on unseen drives.")],
    takeaway="A one-second warning would let the network prepare, or decide not to hand over at all.",
    cite_text="Left: counted from the four measurement campaigns. Right: a review of 22 comparable published models.")

fig_slide(prs, "3 · Literature",
    "Two literatures exist, and neither gives a usable predictor",
    "fig06_protocol_audit",
    [("Measurement studies ", "have exactly the ground truth we need - and build no predictor."),
     ("Prediction studies ", "build models, but mostly on simulation and mostly without a proper split."),
     ("We screened 108 papers; ", "22 of them predict handover, link failure or the next cell."),
     ("Of those 22: ", "none tests on unseen drives, none reports whether its probabilities are trustworthy, and none reports how much warning it gives.")],
    takeaway="The review is about method, not about anyone's numbers: we are saying we cannot tell, not that they are wrong.",
    cite_text="108 references screened over three passes; 22 are comparable models, reviewed on a fixed checklist.")

bullet_slide(prs, "3 · Objectives", "Five objectives",
    [("O1  Predict the next handover ", "at five look-ahead times - 0.5, 1, 2, 3 and 5 seconds - from what a phone can observe at that moment."),
     ("O2  Compare models fairly, ", "under a test that cannot leak information from training into testing."),
     ("O3  Make the probabilities usable: ", "consistent across look-ahead times, correctly calibrated, and with a guaranteed limit on missed handovers."),
     ("O4  Check it still works ", "on a road the model has never seen."),
     ("O5  Explain why it works, ", "and what the network is actually doing.")],
    size=18,
    takeaway="Each objective is answered later in the talk, and the conclusion says how far each one goes.")

# ================================ 4. RESEARCH METHODOLOGY ===================
s = add_slide(prs)
head(s, "4 · Research methodology", "The work follows five simple steps")
tf = textbox(s, L + 0.20, BODY_TOP - 0.02, 11.2, 3.6)
bullets(tf, [
    ("1.  Collect.  ", "Drive a measurement handset over real roads, recording the radio measurements once per second and the full LTE signalling."),
    ("2.  Label.  ", "Read the signalling to find the exact moment of every handover, and mark each sample with whether one follows within 0.5 to 5 seconds."),
    ("3.  Describe.  ", "Build inputs from the past only - signal strength and quality, how fast they are changing, speed and position, and time on the current cell."),
    ("4.  Predict.  ", "Train on some drives and test on completely different drives."),
    ("5.  Check.  ", "Measure not only accuracy, but how trustworthy the probabilities are, how much warning is given, and how many false alarms it costs."),
], size=18, space_after=15)
takeaway_bar(s, "Every number is produced by an automated pipeline from the raw recordings, so the study can be rebuilt from scratch.")

# ==================================== 5. PROPOSED SYSTEM ====================
flow_slide(prs, "5 · Proposed system", "The proposed system, in five stages", "m14_framework",
    [("Stage 02 is what makes the rest possible. ", "A handover is read from the decoded RRC command itself, with a millisecond timestamp."),
     ("Stage 03 uses the past only. ", "107 inputs: radio, mobility, history and signalling."),
     ("Stage 04 is the modelling idea ", "on the next slide."),
     ("Stage 05 sets the alarm threshold ", "on held-out drives, so the miss rate has a guarantee attached.")],
    dh=1.35,
    takeaway="No simulation anywhere: every stage runs on measured data from a live operator network.")

flow_slide(prs, "5 · Proposed system", "Where the ground truth comes from", "m02_signalling",
    [("A handover is a decoded RRC command ", "with a millisecond timestamp - not a counter we have to trust."),
     ("But the configuration arrives in pieces, ", "and its identifiers only mean something against the configuration in force at that moment."),
     ("Our first parse read them flat, and was wrong: ", "it claimed 99.4% of reports were A3, which is impossible, because 43% of reports carry no neighbour at all."),
     ("So we rebuild the configuration timeline ", "and resolve every report against the state actually in force.")],
    takeaway="We found this by checking our extraction against a published method, not by reading our own code.")

flow_slide(prs, "5 · Proposed system", "What the model is given", "m13_features",
    [("152 columns are built and 107 are used, ", "after removing those that are empty or constant."),
     ("Four groups: ", "radio, mobility, history and signalling."),
     ("Everything looks backwards only, ", "and no window crosses from one drive into another."),
     ("One guard matters most: ", "a report arrives 50-200 ms before its handover command - inside one sample - so report counts stop one sample early.")],
    dh=2.30,
    takeaway="The smallest group is the strongest: seven history features, and time on the current cell alone reaches 0.874.")

s = add_slide(prs)
head(s, "5 · Proposed system", "One model, not five")
callout(s, 1.2, 2.38, 10.9, 1.10,
        ["*P(handover within k seconds)  =  1 - (1 - h1)(1 - h2) ... (1 - hk)"], size=21)
tf = textbox(s, L, 3.72, 11.4, 2.3)
bullets(tf, [
    ("The obvious approach fails.  ", "Five separate classifiers, one per look-ahead time, can return a higher probability at 1 second than at 3 seconds - which is impossible."),
    ("So one model predicts h,  ", "the chance the handover happens in the very next second, given that it has not happened yet."),
    ("The five answers are products of h,  ", "so they can only grow with time. Consistency comes from the arithmetic, not a correction step."),
], size=18, space_after=13)
takeaway_bar(s, "This holds for any learner placed inside it - it is a property of the formula, not of the model.")

# ==================== 6. THEORY, MODELLING, SIMULATION, EXPERIMENTS =========
bullet_slide(prs, "6 · Theory and modelling", "What was modelled, and what was compared",
    [("The problem is a time-to-event problem. ", "Instead of asking 'will it happen in the next k seconds', ask 'given it has not happened yet, does it happen now'. That is a hazard."),
     ("Seven models were compared ", "under exactly the same formulation: gradient-boosted trees (LightGBM), logistic regression, a multi-layer perceptron, a temporal convolutional network, a Transformer, a GRU, and the network's own A3 rule as a baseline."),
     ("No simulation was used. ", "Every result in this presentation comes from measured drive-test data on a live network."),
     ("The metric is read against its floor. ", "Handovers are rare, so accuracy would be misleading; every score is compared with what random guessing would achieve at the same rate.")],
    size=18,
    takeaway="The same formulation, the same folds and the same seeds for every model - so the comparison is about the model.")

fig_slide(prs, "6 · Experiments",
    "Where the data was collected",
    "fig03_map_routes",
    [("Four campaigns in Dhaka and Gazipur, ", "57 drives and about 95 km of driving."),
     ("Three urban corridors ", "and one highway run out to Gazipur."),
     ("The circles are handovers. ", "They cluster at particular junctions rather than spreading evenly along the route."),
     ("The drive is the unit of everything: ", "how the data is split, how it is tested, and what the guarantee is attached to.")],
    portrait=True,
    cite_text="GPS from the measurement export. Map data (c) OpenStreetMap contributors, ODbL.")

table_slide(prs, "6 · Experiments",
    "Four measurement campaigns on a live operator network",
    ["Campaign (2026)", "Route", "Drives", "Samples at 1 Hz", "Handovers", "Duration"],
    [["10 September", "urban arterial", "15", "2,700", "290", "46 min"],
     ["12 September", "urban loop", "8", "1,440", "174", "24 min"],
     ["13 September", "dense urban", "20", "3,600", "297", "60 min"],
     ["15 September", "Uttara-Gazipur highway", "14", "2,520", "177", "43 min"],
     ["Total", "two mobility regimes", "57", "10,260", "938", "2.9 h"]],
    col_w=[2.45, 2.95, 1.35, 2.05, 1.45, 1.25], fs=14, hi_last_col=False,
    note="A drive is a continuous segment of at least 60 s. A whole drive is either trained on or tested on - never both.",
    takeaway="The highway run was driven after every modelling decision had been fixed, so it is a genuinely unseen test route.")

fig_slide(prs, "6 · Experiments",
    "Every look-ahead time is a rare-event problem",
    "fig05_dataset",
    [("10,260 samples ", "on a uniform one-second grid."),
     ("The handover rate rises ", "from 3.7% at half a second to 25.9% at five seconds."),
     ("So accuracy would mislead: ", "a model that always answers 'no handover' scores 93.3% at the one-second look-ahead."),
     ("Every score in this talk ", "is therefore read against the rate chance alone would achieve.")],
    takeaway="The one-second grid limits how many events can be detected at all (90.5%), not how far ahead we can predict.")

flow_slide(prs, "6 · Experiments", "How the models were tested", "m03_protocol",
    [("A whole drive is either trained on or tested on ", "- never both."),
     ("Four rotations, ", "so every drive is tested exactly once."),
     ("Repeated with five random seeds, ", "which gives twenty paired comparisons between any two models."),
     ("Confidence intervals ", "come from resampling whole drives, not individual samples.")],
    takeaway="Scalers and alarm thresholds are fitted inside the training folds only, never on the drives used for testing.")

# ========= 7. KNOWLEDGE PROFILE, COMPLEX PROBLEMS AND ACTIVITIES ============
table_slide(prs, "7 · Knowledge profile",
    "Knowledge profile addressed (WK1-WK8)",
    ["", "Attribute", "How this work addresses it"],
    [["WK1", "Natural sciences", "Radio-wave propagation, path loss and shadowing - why signal strength falls at a cell edge."],
     ["WK2", "Mathematics and statistics", "Probability, time-to-event modelling, resampling for confidence intervals, statistical testing."],
     ["WK3", "Engineering fundamentals", "Digital communication, cellular network architecture and mobility management."],
     ["WK4", "Specialist knowledge", "3GPP LTE RRC protocol (TS 36.331): Event A3, measurement reporting, the handover procedure."],
     ["WK5", "Engineering design", "Design of the prediction system, the input feature set and the testing protocol."],
     ["WK6", "Engineering practice", "Drive-test instrumentation, decoding of control-plane logs, a reproducible software pipeline."],
     ["WK7", "Role of engineering in society", "Effects on user experience, network signalling load and handset energy - section 8."],
     ["WK8", "Research literature", "108 research papers screened; 22 comparable models reviewed on a common checklist."]],
    col_w=[0.80, 2.85, 7.85], fs=11.5, hi_last_col=False, body_left=True,
    takeaway="The work needs WK3 to WK8 together; no part of it follows from textbook knowledge alone.")

table_slide(prs, "7 · Complex engineering problem",
    "Complex engineering problem attributes (P1-P7)",
    ["", "Attribute", "How it arises in this work"],
    [["P1", "Depth of knowledge required", "Requires WK3 to WK8 together - protocol, statistics, design and measurement practice."],
     ["P2", "Range of conflicting requirements", "Earlier warning conflicts with more false alarms; a tighter guarantee conflicts with a usable alarm rate."],
     ["P3", "Depth of analysis required", "No standard formulation exists. The time-to-event view was reached after five separate classifiers gave inconsistent answers."],
     ["P4", "Familiarity of issues", "Measurement identifiers are valid only against the configuration in force; overlapping time windows can leak across a careless split."],
     ["P5", "Extent of applicable codes", "3GPP defines the handover procedure, but no standard covers predicting it, or calibrating and evaluating such a system."],
     ["P6", "Extent of stakeholder involvement", "Users, the operator and the regulator want different things: experience, signalling cost and spectrum efficiency."],
     ["P7", "Interdependence", "Signalling decoding, features, modelling, evaluation and uncertainty - a change in one changes the others."]],
    col_w=[0.80, 3.05, 7.65], fs=11.5, hi_last_col=False, body_left=True,
    takeaway="All seven attributes are present, and each trade-off in P2 is measured and reported rather than assumed.")

table_slide(prs, "7 · Complex activities",
    "Complex engineering activity attributes (A1-A5)",
    ["", "Attribute", "How it arises in this work"],
    [["A1", "Range of resources", "A licensed drive-test handset, a survey vehicle, computing resources, and a base of 108 research papers."],
     ["A2", "Level of interaction", "Conflicting technical requirements - lead time, false alarms, calibration and the strength of the guarantee - resolved against one another."],
     ["A3", "Innovation", "Established methods from time-to-event analysis and distribution-free uncertainty were applied to handover prediction for the first time in this form."],
     ["A4", "Consequences for society and the environment", "Less unnecessary signalling and handset energy; better connection quality in dense urban areas - section 8."],
     ["A5", "Familiarity", "Beyond coursework: decoding a live operator's control-plane signalling, and designing an evaluation protocol not found in the literature."]],
    col_w=[0.80, 3.20, 7.50], fs=12.5, hi_last_col=False, body_left=True,
    takeaway="The innovation is in where established methods are pointed, and the protocol that makes the result checkable.")

# ================ 8. IMPACT ON HEALTH, SOCIETY, ENVIRONMENT =================
bullet_slide(prs, "8 · Impact",
    "Impact on health, society, environment and sustainability",
    [("Society.  ", "Dhaka has one of the densest mobile-user populations in the world, and on our routes a handover happens every eleven seconds. A one- to two-second warning lets the network prepare the target cell, shortening the gap a user hears in a call or a video."),
     ("Environment and sustainability.  ", "Every unnecessary handover costs signalling messages, base-station processing and handset transmit energy - and about a quarter of handovers here are unnecessary returns. Removing a share of those saves energy, extends battery life and defers the need for more network hardware."),
     ("Health.  ", "No human subjects, no clinical data and no exposure experiments, so there is no direct health effect. The indirect benefit is reliability: dropped connections matter most for emergency calls and remote-health services."),
     ("Ethics and data.  ", "The recordings contain only the test handset's own measurements and GPS trace. No subscriber data, user traffic or third-party personal information was collected.")],
    size=16,
    takeaway="The ping-pong finding gives the operator one concrete, low-cost configuration change rather than a general recommendation.")

# ============================ 9. RESULTS AND DISCUSSIONS ====================
fig_slide(prs, "9 · Results",
    "The handover is predictable a second before the rule can react",
    "fig10_model_comparison",
    [("At one second: ", "AUROC 0.933, and a precision-recall score of 0.784 against a 6.7% handover rate - nearly twelve times what random guessing achieves."),
     ("The network's own A3 rule, ", "scored on the same data as a predictor, reaches only 0.653."),
     ("Simple models win here. ", "Gradient-boosted trees lead and logistic regression is second, ahead of every deep model - 57 drives is not enough data for them.")],
    takeaway="Every score is read against the prevalence floor, because accuracy on a rare event would be meaningless.",
    cite_text="Out-of-fold results over 57 drives from four campaigns, under the grouped-drive test protocol.")

fig_slide(prs, "9 · Results",
    "How a model is tested can change which model appears to win",
    "fig12_leakage",
    [("Same data, same models. ", "The only change is splitting by whole drive rather than by random sample."),
     ("A GRU inflates by 74%; ", "logistic regression by only 4%."),
     ("Why: ", "these models read a 10-second window, so a randomly split sample shares its own history with the training set."),
     ("The consequence: ", "a careless test does not simply raise all scores - it reorders them.")],
    takeaway="This is a caution for the field, not only for this work: ten of 22 reviewed papers state no split protocol at all.")

fig_slide(prs, "9 · Results",
    "Consistent across look-ahead times, and calibrated",
    "fig13_hazard_results",
    [("Five separate classifiers contradict themselves ", "on 43.6% of samples - a higher chance at 1 second than at 3."),
     ("The single-model formulation contradicts itself on 0%, ", "by construction."),
     ("Calibration improves too, ", "at every look-ahead time, against an uncalibrated baseline."),
     ("Honest point: ", "a baseline with its own calibration step can beat us on calibration alone - but it costs held-out drives and 2 to 5 points of accuracy.")],
    takeaway="The claim is narrow and defensible: consistency, calibration and ranking together, from a single fit.")

fig_slide(prs, "9 · Results",
    "The warning threshold comes with a guarantee, and a price",
    "fig14_riskcontrol",
    [("The threshold is not chosen by hand. ", "It is calibrated on 28 drives held out for the purpose."),
     ("At a 20% target miss rate: ", "it alarms on 23% of samples and actually misses 12.6% of handovers - an operating point a network could run."),
     ("At a 5% target: ", "the alarm rate is 61%, which nobody would deploy."),
     ("How many drives you calibrate on ", "sets the tightest target you can even ask for - which makes it something to plan a campaign around.")],
    takeaway="The whole trade-off curve is reported, not the single point that flatters the result.")

fig_slide(prs, "9 · Results",
    "It still works on a road the model has never seen",
    "fig25_capture_transfer",
    [("Each campaign held out completely: ", "AUROC stays between 0.909 and 0.949."),
     ("The highway scores 0.927, ", "with the best calibration of all four - a corridor and a speed the model had never seen."),
     ("It was driven after ", "every modelling decision had been fixed, so nothing was tuned for it."),
     ("Honest limit: ", "four campaigns give four test points. A strong design, not a large sample.")],
    takeaway="Adding the fourth campaign changed the headline result by nothing at all: 0.933 before, 0.933 after.")

fig_slide(prs, "9 · Results",
    "It also works on another team's dataset",
    "fig15_transfer",
    [("Between our own campaigns: ", "AUROC stays between 0.88 and 0.93, with nothing shared - no drives, no routes, no cells."),
     ("On a public dataset ", "collected by a different group, we reach 0.752."),
     ("A model trained on that data ", "reaches 0.745 on itself."),
     ("The absolute level is lower ", "because that dataset has no GPS, so our mobility inputs are unavailable.")],
    takeaway="Not degrading - reaching what that data supports.",
    cite_text="Public dataset: Shafi et al. (2025), Mendeley Data, DOI 10.17632/n2pvmtyn2j.1.")

fig_slide(prs, "9 · Discussion",
    "Why it works: the network triggers on its weakest signal",
    "fig17_mechanism",
    [("The best single input ", "is how long the phone has already been on this cell: 0.874 on its own."),
     ("The quantity the rule uses ", "- the gap between serving and neighbour - scores only 0.566."),
     ("The reason: ", "the gap says a handover is allowed. The dwell time says one is due."),
     ("It is also free: ", "dwell time is a counter the phone already keeps.")],
    takeaway="This is the finding an operator can act on, and it needs no new measurement to obtain.")

flow_slide(prs, "9 · Discussion", "Three in five A3 reports are never acted on", "m09_conversion",
    [("7,385 A3 reports ", "were sent across the four campaigns."),
     ("62.9% of them ", "are not followed by a handover within two seconds."),
     ("And the rate rises with cell size: ", "57% on the urban arterial, 63% in the dense core, 72% on the highway."),
     ("Reporting a condition ", "is not the same as acting on it.")],
    takeaway="Two qualifiers always travel with this number: which reports were counted, and over which window.")

fig_slide(prs, "9 · Discussion",
    "Where ping-pong actually comes from",
    "fig24_pingpong_mechanism",
    [("Handovers that stay on the same carrier ", "return at 31.0%; those that change carrier, at 6.5%."),
     ("One configuration carries it: ", "a +1 dB offset with a 320 ms time-to-trigger, which covers 72% of all handovers."),
     ("It is not speed. ", "The highway returns 31.1% - essentially the same as the urban rate."),
     ("The lever is that time-to-trigger, ", "and it is a parameter an operator can change.")],
    takeaway="This is the one concrete configuration recommendation the thesis can make.")

fig_slide(prs, "9 · Discussion",
    "What did not help, and what did",
    "fig22_negatives",
    [("Added in good faith, and measured: ", "more neighbour data, explicit clustering features, a dedicated ping-pong model, and two domain-adaptation methods."),
     ("None of them helped. ", "The dedicated ping-pong model came out at chance; domain adaptation made transfer worse."),
     ("What did move the result: ", "the formulation, the configuration regime, and splitting the data by whole drive."),
     ("Reported rather than dropped, ", "because a pattern across five independent attempts is itself a result.")],
    takeaway="Every negative here is a controlled comparison under the same protocol, not an absence of effort.")

# ==================================== 10. CONCLUSION ========================
s = add_slide(prs)
head(s, "10 · Conclusion", "Conclusion")
tf = textbox(s, L, BODY_TOP + 0.05, 11.4, 3.5)
bullets(tf, [
    ("1.  The next handover can be predicted about a second before the network's own rule can react,  ",
     "using only what a handset already measures - AUROC 0.933, with calibrated probabilities and a guaranteed limit on missed handovers."),
    ("2.  How the problem is formulated matters more than adding more data.  ",
     "One model of the next second beats five separate classifiers, and gives consistent answers for free."),
    ("3.  How a model is tested is part of the contribution.  ",
     "A careless split inflates a deep model by 74% and a linear one by 4%, which changes the ranking itself."),
    ("4.  It generalises to an unseen road,  ",
     "scoring 0.927 on a highway corridor driven after every modelling decision was frozen."),
], size=17, space_after=15)
takeaway_bar(s, "Stated plainly: this is an offline predictor with a measured guarantee - not yet a deployed handover controller.")

# =================================== 11. FUTURE WORKS =======================
bullet_slide(prs, "11 · Future works", "Future works",
    [("Collect more routes and more operators.  ", "Four days on one network is enough for a careful study, not enough to claim the result holds everywhere."),
     ("Measure the actual benefit.  ", "Link the warning to a measured throughput or interruption improvement, rather than the counting estimate reported here."),
     ("Test one configuration change with the operator.  ", "A longer time-to-trigger on the layer where ping-pong concentrates is a single, low-cost experiment."),
     ("Release the code and the dataset.  ", "Only two of the 22 papers reviewed release code and only one releases data.")],
    size=18,
    takeaway="The next milestone is a field trial: the same warning, measured against a real network outcome.")

# ---------------------------------------------------------------- closing
s = add_slide(prs, BG_TITLE)
tf = textbox(s, 0.80, 1.65, 8.4, 1.2)
para(tf, "Thank you", 44, MAROON, bold=True, first=True, space_after=0)
tf = textbox(s, 0.80, 2.95, 8.4, 1.4)
para(tf, "Questions and discussion", 22, INK, first=True, space_after=12)
para(tf, "Abeer Saadman  |  Department of EEE, Islamic University of Technology", 15, MUTED)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("slides:", len(prs.slides._sldIdLst))
print("saved", OUT)
