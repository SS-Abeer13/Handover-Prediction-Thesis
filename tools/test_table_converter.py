import docx, re

doc = docx.Document('Docs/Handover_Thesis_Manuscript_Revised.docx')
print(f"Loaded {len(doc.tables)} tables.")

def clean_cell(text):
    text = text.strip().replace('\n', ' ')
    # Date replacements
    text = re.sub(r'10\s*Sept(?:ember)?\s*\(([^)]+)\)', r'1st Campaign (\1)', text, flags=re.I)
    text = re.sub(r'12\s*Sept(?:ember)?\s*\(([^)]+)\)', r'2nd Campaign (\1)', text, flags=re.I)
    text = re.sub(r'13\s*Sept(?:ember)?\s*\(([^)]+)\)', r'3rd Campaign (\1)', text, flags=re.I)
    text = re.sub(r'15\s*Sept(?:ember)?\s*\(([^)]+)\)', r'4th Campaign (\1)', text, flags=re.I)
    text = re.sub(r'\b10\s*(?:th\s*(?:of\s*)?)?Sept(?:ember)?\b', '1st Campaign', text, flags=re.I)
    text = re.sub(r'\b12\s*(?:th\s*(?:of\s*)?)?Sept(?:ember)?\b', '2nd Campaign', text, flags=re.I)
    text = re.sub(r'\b13\s*(?:th\s*(?:of\s*)?)?Sept(?:ember)?\b', '3rd Campaign', text, flags=re.I)
    text = re.sub(r'\b15\s*(?:th\s*(?:of\s*)?)?Sept(?:ember)?\b', '4th Campaign', text, flags=re.I)
    text = re.sub(r'Campaign\s*\(2026\)', 'Campaign', text, flags=re.I)
    # Unicode replacements
    text = text.replace('%', r'\%')
    text = text.replace('&', r'\&')
    text = text.replace('_', r'\_')
    text = text.replace('—', '---').replace('–', '--')
    text = text.replace('±', r'$\pm$').replace('²', r'$^2$').replace('³', r'$^3$')
    text = text.replace('α', r'$\alpha$').replace('β', r'$\beta$').replace('δ', r'$\delta$')
    text = text.replace('λ', r'$\lambda$').replace('μ', r'$\mu$').replace('τ', r'$\tau$')
    text = text.replace('≤', r'$\le$').replace('≥', r'$\ge$').replace('≈', r'$\approx$')
    text = text.replace('×', r'$\times$').replace('·', r'$\cdot$')
    text = text.replace('“', "``").replace('”', "''").replace('’', "'").replace('‘', "`")
    return text

for i in range(1, len(doc.tables)):
    tbl = doc.tables[i]
    num_cols = len(tbl.columns)
    header = [clean_cell(c.text) for c in tbl.rows[0].cells]
    print(f"Table {i:2d}: {len(tbl.rows)} rows x {num_cols} cols | {header[0][:30]}")
