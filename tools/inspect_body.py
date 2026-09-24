import docx

doc = docx.Document('Docs/Handover_Thesis_Manuscript_Revised.docx')
body = doc._element.body

p_idx = 0
t_idx = 0
elements = []
for child in body:
    if child.tag.endswith('p'):
        elements.append(('p', p_idx))
        p_idx += 1
    elif child.tag.endswith('tbl'):
        elements.append(('tbl', t_idx))
        t_idx += 1

print(f"Total paragraphs: {p_idx}, Total tables: {t_idx}, Total elements: {len(elements)}")

for pos, (elem_type, idx) in enumerate(elements):
    if elem_type == 'tbl':
        prev_p = [e[1] for e in elements[:pos] if e[0] == 'p']
        last_p_idx = prev_p[-1] if prev_p else None
        last_p_txt = doc.paragraphs[last_p_idx].text.strip()[:60] if last_p_idx is not None else ''
        tbl = doc.tables[idx]
        first_cell = tbl.cell(0, 0).text.strip().replace('\n', ' ')[:30]
        safe_p = last_p_txt.encode('ascii', 'replace').decode()
        safe_c = first_cell.encode('ascii', 'replace').decode()
        print(f"Table {idx:2d} after P{last_p_idx:3d} [{safe_p}] -> {safe_c}")
