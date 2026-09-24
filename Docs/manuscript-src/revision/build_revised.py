# -*- coding: utf-8 -*-
import sys
from pathlib import Path
import copy
import shutil
import docx
from docx.oxml.ns import qn
from docx.shared import Emu, Inches
from PIL import Image

sys.path.insert(0, str(Path("d:/Handover Thesis/Docs/manuscript-src/revision").resolve()))
from msnumbers import build
from refs import REFS
from revtext_a import DELETE, TITLE, part_a
from revtext_b import part_b
from revtext_c import part_c
from tables_spec import appendix_d, tables

ROOT = Path("d:/Handover Thesis").resolve()
SRC = ROOT / "Docs" / "Handover_Thesis_Manuscript_Draft.docx"
FIGS = ROOT / "pipeline" / "reports_rev" / "figures"
OUT = ROOT / "Claude outputs" / "Handover_Thesis_Manuscript_Revised-5.docx"
MIRROR = ROOT / "Docs" / "Handover_Thesis_Manuscript_Revised.docx"

TABLE_SLOT = {1: "2.1", 2: "3.1", 3: "3.2", 4: "4.1", 5: "4.2", 6: "4.3", 7: "5.1", 8: "5.3", 9: "5.5",
              10: "5.6", 11: "5.7", 12: "5.8", 13: "5.9", 14: "5.10", 15: "5.12", 16: "5.13", 17: "5.14",
              18: "5.15", 25: "A.1", 26: "A.2"}

TABLE_CAPTION_PARA = {1: 277, 2: 350, 3: 359, 4: 381, 5: 410, 6: 422, 7: 477, 8: 492, 9: 502, 10: 510,
                      11: 525, 12: 537, 13: 545, 14: 553, 15: 565, 16: 574, 17: 581, 18: 593,
                      25: 729, 26: 731}

TABLE_CAPTIONS = {
 "2.1": "Table 2.1 Protocol audit of the 22 most comparable prediction models, contrasted with the protocol "
        "adopted here. Appendix D lists the models.",
 "3.1": "Table 3.1 Deployed Event A3 configuration recovered from RRC signalling, weighted by the handovers each "
        "profile produced.",
 "3.2": "Table 3.2 The four measurement campaigns after quality control.",
 "4.1": "Table 4.1 Positive-class prevalence by horizon on the common row set, and the accuracy obtained by a "
        "constant negative prediction.",
 "4.2": "Table 4.2 Measured discrete-time hazard by bin, and the cumulative incidence it implies, against the "
        "prevalence of Table 4.1. The two agree exactly by construction.",
 "4.3": "Table 4.3 Feature blocks, their source and their alignment.",
 "5.1": "Table 5.1 The alignment audit: share of handovers whose target cell already serves in export rows around "
        "the command.",
 "5.2": "Table 5.2 What the one-row lag costs, with everything else held fixed.",
 "5.3": "Table 5.3 Out-of-fold performance of LightGBM under the hazard formulation, leave-one-campaign-out. "
        "ARM A (tuned): nested random search of twenty trials per outer fold, averaged over three seeds. Tables 5.3, 5.4, 5.8 and A.1 are this arm, and every headline number in the abstract and Chapter 7 comes from it. "
        "Intervals are block-bootstrapped over 180 s blocks and are conditional on these four sessions.",
 "5.4": "Table 5.4 Every learner under identical folds, features, formulation and tuning budget. "
        "ARM A (tuned): nested random search of twenty trials per outer fold, averaged over three seeds. Tables 5.3, 5.4, 5.8 and A.1 are this arm, and every headline number in the abstract and Chapter 7 comes from it.",
 "5.5": "Table 5.5 One-second AUPRC under four splitting protocols. ARM B (fixed hyper-parameters): "
        "one configuration per learner, held fixed across all four protocols so that the splitting rule "
        "is the only thing that varies. Because the learners are untuned, the levels in the "
        "leave-one-campaign-out column are NOT comparable with Tables 5.3, 5.4 and A.1, which report "
        "arm A; every percentage quoted from this table is a ratio within the table.",
 "5.6": "Table 5.6 Horizon coherence and calibration for the hazard formulation and its four controls. "
        "ARM C (coherence): fixed hyper-parameters, with an extra calibration block held out of every "
        "fold so the isotonic arms can be fitted; its AUPRC levels are therefore not comparable with "
        "arm A either, and only the comparison between arms within this table is meaningful. "
        "Violation rates are given for a single fitted system, which is what a deployment meets, and for "
        "the three-seed ensemble beside it; averaging predictions over seeds lowers the rate without "
        "removing its cause. Every violation figure quoted in the text is the single-fit rate.",
 "5.7": "Table 5.7 Conformal risk control at the one-second horizon across regimes. Grouped columns distinguish "
        "exchangeable block pooling (n ≈ 43 calibration units) from held-out cross-campaign deployment "
        "(n ∈ {8, 14, 15, 20} calibration units, mean n = 14.25). Target α = 0.05 is feasible in 25% of "
        "rotations (Campaign 3, n = 20, floor 0.048 ≤ 0.05) where λ̂ = 0.0 achieves 0.000 loss at 100% alarm; "
        "the arithmetic mean floor is 0.072. Em dashes denote infeasible event-level targets bounded by the 36.4% "
        "unresolvable handover floor.",
 "5.8": "Table 5.8 Leave-one-campaign-out performance at the one-second horizon.",
 "5.9": "Table 5.9 Transfer between our campaigns and an independently collected public dataset on the same "
        "network, both built from L3 signalling alone.",
 "5.10": "Table 5.10 Single-feature AUROC at the one-second horizon on the lagged export. The "
         "serving-to-neighbour gap is defined as the serving cell's level minus the best "
         "neighbour's, in decibels, so a smaller gap means the neighbour is closer to overtaking "
         "and the A3 entry condition is closer to being met.",
 "5.11": "Table 5.11 Performance inside and outside handover bursts.",
 "5.12": "Table 5.12 Event A3 report conversion, counted per transmitted report and per trigger episode.",
 "5.13": "Table 5.13 The ping-pong rate on one fixed set of handover commands under three definitions.",
 "5.14": "Table 5.14 Ping-pong rate by carrier relationship and mobility regime.",
 "5.15": "Table 5.15 Approaches implemented and measured that produced no improvement.",
 "A.1": "Table A.1 Per-horizon AUPRC for every learner under leave-one-campaign-out, against the prevalence floor.",
 "A.2": "Table A.2 Event-level operating characteristics of the primary model at the 5 % false-positive "
        "operating point.",
 "5.2b": "Table 5.2b Where the leak actually lives. One common row set of usable rows, identical "
         "folds, seeds and learner; only the set of columns shifted by one row changes. The "
         "serving-cell radio scalars are measured on whichever cell the export treats as serving, so "
         "after the end-of-second flip they describe the target rather than the source.",
 "5.6b": "Table 5.6b Paired differences between the hazard arm and each control at one second. A "
          "positive ECE difference favours the control, since a lower calibration error is better. "
          "The block column is a descriptive within-campaign spread and supports no inference: "
          "Section 5.3 shows 180-second blocks of one session are dependent, so resampling them "
          "understates the uncertainty. The four per-campaign columns are the independent units; "
          "with four of them a two-sided sign test cannot return a p-value below 0.125, so the "
          "count of campaigns sharing the sign is reported and no significance is claimed.",
 "5.16": "Table 5.16 Nine cases drawn at random, with a fixed seed, from the out-of-fold "
         "predictions at the 5 % false-positive operating point: three the model called correctly, "
         "three it alarmed on with no command following, and three it missed. The draw is a uniform "
         "sample within each class, not a selection.",
 "A.5": "Table A.5 Sensitivity of every window chosen after the data were collected. Each is varied "
        "over a range bracketing the value used in Chapter 5, with everything else held fixed. The "
        "block-length and purge rows use fixed hyper-parameters (arm B), so their levels are "
        "comparable with each other and with Table 5.5, not with Table 5.3.",
 "A.6": "Table A.6 Per-handover evidence for the delayed-refresh explanation of the alignment "
        "residual in Section 5.1. The explanation is a claim about individual handovers, so it is "
        "settled one handover at a time rather than asserted.",
 "A.4": "Table A.4 Per-campaign calibration at the one-second horizon for every coherence arm, on the "
        "row set and hyper-parameters of Table 5.6 (arm C). Section 5.4 refers to these numbers; an "
        "earlier draft promised them here and printed a per-campaign count taken from the tuned arm "
        "instead, which is a different experiment.",
 "3.3": "Table 3.3 Ground-truth integrity. Every handover command in the event set carries a decoded "
        "RRCConnectionReconfigurationComplete, so the label is a completed handover and not merely an "
        "issued command; the re-establishments recorded alongside are largely not consequent on it.",
 "A.3": "Table A.3 The alignment audit of Section 5.1 across three datasets and two instruments. The "
        "defect follows the row-aggregation rule, not the vendor and not the network: it is present "
        "wherever a row summarises the second that ends at its timestamp, absent where rows are "
        "instantaneous samples, and it disappears from the same NUWiNS handovers when the "
        "downsampling rule is changed to the start of the second.",
 "D.1": "Table D.1 The 22 audited prediction models.",
}

INSERT_AFTER_TABLE = {"3.3": 3, "5.2b": 7, "5.2": 7, "5.4": 8, "5.6b": 10, "5.11": 14,
                      "A.6": 26, "A.5": 26, "A.4": 26, "A.3": 26, "5.16": 18}

FIGURES = {
 219: (None, "Figure 1.1 Event A3 fires only after the neighbour has been better than the serving cell by a "
             "configured offset for a full time-to-trigger. The profile that produces most of the handovers "
             "measured here is a +1 dB offset with a 320 ms time-to-trigger and 1 dB hysteresis."),
 304: (None, "Figure 2.1 The methodology in five stages. Stage 2 supplies the ground truth; Stage 4 is the "
             "formulation that keeps the five horizons mutually consistent; Stage 5 is the evaluation protocol "
             "whose unit is a whole campaign."),
 329: (None, "Figure 3.1 The three Dhaka campaigns on an OpenStreetMap background. Events cluster at particular "
             "junctions rather than distributing uniformly along the route. The 15 September highway campaign "
             "lies outside this extent. Map data © OpenStreetMap contributors."),
 341: (None, "Figure 3.2 Measurement identifiers persist in VarMeasConfig until they are modified or removed. A "
             "flat parse resolves every report against the union of all configuration fragments; the timeline "
             "replays additions, modifications and removals in order and resolves each report against the "
             "configuration in force at its own timestamp."),
 355: ("fig_c5_profiles.png",
       "Figure 3.3 Handovers attributed to each configuration profile, weighted by the handovers each actually "
       "produced rather than by the number of configurations present. Blue bars are positive-offset A3 "
       "profiles, red bars negative-offset A3 profiles, grey bars other events."),
 362: (None, "Figure 3.4 Composition of the pooled dataset. The four campaigns differ in duration and in event "
             "density; the highway campaign contributes the fewest handovers per minute and the highest mean "
             "speed."),
 403: ("fig_c11_schematic.png",
       "Figure 4.1 Independent per-horizon classifiers against the hazard formulation, drawn on simulated rows "
       "with the measured violation rate. Independent fits cross; the product of hazards cannot."),
 444: ("fig_c1_protocols.png",
       "Figure 4.2 The splitting ladder. Leave-one-campaign-out holds out a whole continuous session; blocked "
       "cross-validation holds out a contiguous stretch of each session with a 60-second purge on either side; "
       "the two leaky protocols draw 180-second blocks, or rows, at random."),
 484: ("fig_c9_alignment.png",
       "Figure 5.1 The alignment audit and its mechanism. (a) A handover command at tau inside the second [t, t+1) "
       "is carried by the row stamped t whenever the serving-cell register updates before the second closes, that "
       "is whenever the margin m = 1 - delta - x is positive. (b) The audit on our captures: rows around the "
       "command. (c) The same quantity on three datasets; the two aggregated exports are contaminated and the "
       "sampled log is not. (d) The falsification test: no handover whose execution crosses the second boundary "
       "contaminates its row. (e) Contamination against where the command falls inside the second, in both XCAL "
       "corpora. (f) In the last fifth of the second the contamination share becomes an estimator of the "
       "serving-cell update delay, which sits near 60 ms."),
 486: ("fig_c18_events.png",
       "Figure 5.2 Event-level cost of the warning at the 5 % false-positive operating point: share of handover "
       "commands detected against the 1 Hz grid ceiling, with the false-alarm episode rate per hour on the right "
       "axis."),
 490: ("fig_c3_main.png",
       "Figure 5.3 AUPRC lift over the prevalence floor by horizon for every learner, out of fold under "
       "leave-one-campaign-out. A lift of one is the floor itself."),
 497: ("fig_c3_per_capture.png",
       "Figure 5.4 One-second AUPRC lift by learner, with each held-out campaign drawn separately. The spread "
       "across campaigns is wider than the spread across learners."),
 504: ("fig_c2_protocols.png",
       "Figure 5.5 The splitting ladder, at one and five seconds. Every learner rises as the protocol weakens, "
       "and the learners rise by different amounts, so the ranking changes with the protocol."),
 516: ("fig_c11_paired.png",
       "Figure 5.6 Calibration error at one second, one line per held-out campaign, for independent classifiers "
       "against the hazard formulation."),
 518: ("fig_c11_coherence.png",
       "Figure 5.7 Calibration error against the horizon for every coherence arm (left) and the share of rows "
       "violating horizon ordering (right)."),
 520: ("fig_c8_bursts.png",
       "Figure 5.8 Where the model earns its advantage: rows with no handover in the previous ten seconds "
       "against rows inside a burst."),
 532: ("fig_c10_crc.png",
       "Figure 5.9 Conformal risk control in both regimes. Left: realised mean per-unit miss rate against the "
       "target, with the guarantee boundary as the diagonal. Right: the alarm rate that compliance costs."),
 541: ("fig_c3_loco.png",
       "Figure 5.10 Leave-one-campaign-out on four metrics. Each bar holds one campaign out entirely and trains "
       "on the other three; the dashed line is the pooled out-of-fold value. The highway campaign is in red."),
 548: ("fig_c6_external.png",
       "Figure 5.11 Transfer to an independently collected dataset on the same network, with both domains built "
       "from L3 signalling alone. Bars are 95 % bootstrap intervals over runs."),
 555: ("fig_c14_single.png",
       "Figure 5.12 Single-feature discriminative power at one second on the lagged export. No single feature "
       "approaches the full model."),
 561: ("fig_c8_dwell.png",
       "Figure 5.13 Probability of a handover command within one and five seconds against the time since the "
       "previous command, computed from signalling timestamps."),
 567: ("fig_c7_conversion.png",
       "Figure 5.14 A3 report conversion under three units. Counting transmitted reports counts periodic "
       "re-reports of one trigger; counting trigger episodes loosely lets several overlapping episodes "
       "claim the same command; the one-to-one unit matches each command to exactly one episode."),
 578: ("fig_c13_grid.png",
       "Figure 5.15 The ping-pong rate over the full definitional grid: three return windows against three "
       "conventions for cell identity and return, computed on one unchanging set of handover commands."),
 583: ("fig_c13_strata.png",
       "Figure 5.16 Where ping-pong concentrates. The dominant axis is the carrier relationship, not the "
       "mobility regime."),
 587: (DELETE, None),
 589: ("fig_c_hawkes_branching.png",
       "Figure 5.17 Hawkes branching ratio estimated on each campaign separately and on the pooled stream, "
       "with bootstrap intervals."),
 595: ("fig_c_hawkes_residuals.png",
       "Figure 5.18 Self-excitation in the handover arrival stream. The clustering is real; the Ogata "
       "residual test rejects the exponential kernel, so its temporal shape is not exponential."),
 600: ("fig_c18_ablation.png",
       "Figure 5.19 Feature-block, alignment and censoring ablations at one and five seconds, out of fold under "
       "leave-one-campaign-out, with the spread over three seeds."),
}

NEW_FIG_CAPTION = (
    "Figure 5.20 Nine cases drawn uniformly at random, with a fixed seed, from the out-of-fold "
    "predictions at the 5 % false-positive operating point. Blue is the serving cell's RSRP, red the "
    "best neighbour's, green the model's probability of a command within one second against the "
    "right axis, the dotted line the alarm threshold, black lines are handover commands and the "
    "orange line is the drawn row. Nothing here was selected for being illustrative.")

def set_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    lead_bold = bool(p.runs[0].bold)
    mixed = lead_bold and any(not r.bold for r in p.runs[1:] if r.text.strip())
    head, sep, tail = text.partition(". ")
    if mixed and sep and len(head) < 60:
        p.runs[0].text = head + sep
        if len(p.runs) > 1:
            p.runs[1].text = tail
            p.runs[1].bold = False
            for r in p.runs[2:]:
                r.text = ""
        else:
            r = p.add_run(tail)
            r.bold = False
        return
    p.runs[0].text = text
    if mixed:
        p.runs[0].bold = False
    for r in p.runs[1:]:
        r.text = ""

def insert_after(p, texts, style=None):
    prev = p
    made = []
    for txt in texts:
        new = copy.deepcopy(p._p)
        prev._p.addnext(new)
        np_ = docx.text.paragraph.Paragraph(new, p._parent)
        for el in np_._p.findall(docx.oxml.ns.qn("w:hyperlink")):
            np_._p.remove(el)
        set_text(np_, txt)
        if style:
            np_.style = style
        prev = np_
        made.append(np_)
    return made

def delete(p):
    p._p.getparent().remove(p._p)

def autofit(tbl):
    tblPr = tbl._tbl.tblPr
    for tag in ("w:tblW", "w:tblLayout"):
        for el in tblPr.findall(qn(tag)):
            tblPr.remove(el)
    layout = tblPr.makeelement(qn("w:tblLayout"), {qn("w:type"): "autofit"})
    tblPr.append(layout)
    for grid in tbl._tbl.findall(qn("w:tblGrid")):
        tbl._tbl.remove(grid)
    g = tbl._tbl.makeelement(qn("w:tblGrid"), {})
    ncol = len(tbl.rows[0].cells)
    for _ in range(ncol):
        g.append(g.makeelement(qn("w:gridCol"), {}))
    tbl._tbl.insert(list(tbl._tbl).index(tblPr) + 1, g)
    for row in tbl.rows:
        for c in row.cells:
            tcPr = c._tc.get_or_add_tcPr()
            for el in tcPr.findall(qn("w:tcW")):
                tcPr.remove(el)

def fill_table(tbl, rows, doc):
    while len(tbl.rows) > 1:
        tbl._tbl.remove(tbl.rows[-1]._tr)
    head = tbl.rows[0]
    ncol = len(head.cells)
    proto = copy.deepcopy(head._tr)
    def write(tr_row, values):
        cells = tr_row.cells
        for i in range(len(cells)):
            v = values[i] if i < len(values) else ""
            cell = cells[i]
            cell.text = ""
            para = cell.paragraphs[0]
            set_text(para, str(v))
    if ncol != len(rows[0]):
        tbl._tbl.remove(head._tr)
        for r in rows:
            tr = copy.deepcopy(proto)
            tbl._tbl.append(tr)
            while len(docx.table._Row(tr, tbl).cells) > len(r):
                tr.remove(tr.tc_lst[-1])
            while len(docx.table._Row(tr, tbl).cells) < len(r):
                tr.append(copy.deepcopy(tr.tc_lst[-1]))
            write(docx.table._Row(tr, tbl), r)
        autofit(tbl)
        return
    write(head, rows[0])
    for r in rows[1:]:
        tr = copy.deepcopy(proto)
        tbl._tbl.append(tr)
        write(docx.table._Row(tr, tbl), r)
    autofit(tbl)

def set_borders(tbl):
    from docx.oxml import OxmlElement
    pr = tbl._tbl.find(qn("w:tblPr"))
    if pr is None:
        pr = OxmlElement("w:tblPr")
        tbl._tbl.insert(0, pr)
    old = pr.find(qn("w:tblBorders"))
    if old is not None:
        pr.remove(old)
    bd = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        bd.append(e)
    pr.append(bd)

def new_table(doc, rows, after_para):
    tbl = doc.add_table(rows=1, cols=len(rows[0]))
    try:
        tbl.style = doc.tables[0].style
    except Exception:
        pass
    set_borders(tbl)
    after_para._p.addnext(tbl._tbl)
    fill_table(tbl, rows, doc)
    return tbl

def replace_image(doc, para, png: Path):
    rid = para._p.xpath(".//a:blip/@r:embed")[0]
    part = doc.part.related_parts[rid]
    data = png.read_bytes()
    part._blob = data
    w, h = Image.open(png).size
    ext = para._p.xpath(".//wp:extent")[0]
    cx = int(ext.get("cx"))
    ext.set("cy", str(int(cx * h / w)))
    for e in para._p.xpath(".//a:ext"):
        e.set("cx", str(cx))
        e.set("cy", str(int(cx * h / w)))

def line_map(doc):
    out = {}
    line = 1
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip():
            out[line] = idx
            line += 1
    return out

def fixups(N):
    R = {k[4:]: N[k] for k in N if k.startswith("ref_")}
    return [
        ("Hawkes [21] is the standard model", f"Hawkes [{R['hawkes']}] is the standard model"),
        ("Laub et al. [22] provide", f"Laub et al. [{R['laub']}] provide"),
        ("residual analysis of Ogata [23], in which",
         f"residual analysis of Ogata [{R['ogata']}], in which"),
        ("Price-Williams and Heard [24] establish",
         f"Price-Williams and Heard [{R['pricew']}] establish"),
        ("Ismail Fawaz et al. [25] benchmark nine unsupervised domain-adaptation algorithms across",
         f"Ismail Fawaz et al. [{R['fawaz']}] benchmark nine unsupervised domain-adaptation "
         "algorithms across"),
        ("Deep CORAL [26] in Chapter 5", f"Deep CORAL [{R['coral']}] in Chapter 5"),
        ("Shi et al. [27] supply", f"Shi et al. [{R['shi']}] supply"),
        ("Wagner et al. [28] show formally", f"Wagner et al. [{R['wagner']}] show formally"),
        ("Deb et al. [29] model conditional handover",
         f"Deb et al. [{R['deb']}] model conditional handover"),
        ("Temperature scaling [31] fits", f"Temperature scaling [{R['guo']}] fits"),
        ("self-exciting Hawkes process [21], [22] with conditional intensity",
         f"self-exciting Hawkes process [{R['hawkes']}], [{R['laub']}] with conditional intensity"),
        ("Price-Williams and Heard [24], the fit is validated by the residual analysis of Ogata [23]",
         f"Price-Williams and Heard [{R['pricew']}], the fit is validated by the residual "
         f"analysis of Ogata [{R['ogata']}]"),

        ("exp( −β ( t − t_i ) )\t(4.12)", "exp( −β ( t − t_i ) )\t(4.13)"),

        ("the row filter of Section 3.7", "the row filter of Section 4.1", 2),
        ("The ping-pong behaviour of Section 5.9 suggests",
         "The ping-pong behaviour of Section 5.10 suggests"),

        ("The 107 features retained after the degenerate-feature filter of Section 4.6 are "
         "organised as follows.",
         f"The {N['n_features']} features of the main set defined in Section 4.6 are organised as "
         "follows. No feature is filtered out: three of them (the missingness indicators for "
         "serving RSRP, RSRQ and SINR) are constant on this corpus, because those fields are never "
         "missing here, and a tree or a regularised linear model ignores a constant column without "
         "being told to. They are retained so that the feature set is defined by its construction "
         "rule rather than by this dataset."),

        # Appendix B feature blocks
        ("Radio block (83 features).", f"Radio block ({N['n_feat_radio']} features)."),
        ("Signalling block (13 features).", f"Signalling block ({N['n_feat_signalling']} features, evaluated in ablation)."),
        ("History block (7 features).", f"History block ({N['n_feat_history']} features)."),

        ("A warning one second ahead of a handover command is sufficient time for the network to "
         "complete target-cell preparation, and a warning that a handover is likely to be reversed "
         "within fifteen seconds is sufficient information to consider not performing it.",
         "A warning far enough ahead of a handover command would leave time for target-cell "
         "preparation, and a warning that a handover is likely to be reversed within fifteen "
         "seconds would be enough to consider not performing it. Both are stated as conditionals "
         "because this thesis does not establish either: Section 5.5 and Appendix A measure what a "
         "warning built on this instrument actually delivers at the one-second horizon — "
         f"{N['det1']:.0%} of commands detected, {N['fa_hour1']:.0f} false-alarm episodes per hour "
         f"and a median lead of {N['lead1']:.2f} s — and an alarm every "
         f"{3600 / N['fa_hour1']:.0f} seconds with under half a second of warning is not an "
         "operating point a radio resource manager can use. The motivation for the task is that "
         "the reactive rule leaves a measurable cost on the table; the finding is that one-hertz "
         "handset telemetry does not recover enough of it."),

        ("The operational meaning of an expected calibration error of 0.024 at the one-second "
         "horizon is worth stating plainly, because a calibration figure is easy to report and easy "
         "to ignore.",
         "Three calibration errors appear in this chapter and they are not interchangeable, so the "
         "arm is named each time from here on: "
         f"{N['ece1']:.3f} is the tuned primary model of Table 5.3 (arm A), "
         f"{N['ece1_hazard']:.3f} is the hazard arm of Table 5.6 (arm C, fitted with an extra "
         f"calibration block held out of every fold) and {N['ece1_isopav']:.3f} is the isotonic "
         "composite in the same table. The operational meaning of a calibration error at the "
         "one-second horizon is worth stating plainly, because such a figure is easy to report and "
         "easy to ignore."),

        ("Alongside the {ho} commands, the signalling log records".format(ho=N["ho_all"]),
         f"Table 3.3 also reconciles the two re-establishment counts, which do not match by "
         f"design: Table 3.2 counts every re-establishment in each session, while Table 3.3 counts "
         "only those close to a handover command. The difference is the point — most "
         "re-establishments in this corpus are not consequent on mobility at all, and the "
         f"{N['reest_all'] - N['gt_reest_near']} that fall outside the two-second window after a "
         "command occur during ordinary connected-mode operation, predominantly at session start "
         "and after coverage gaps. Alongside the "
         f"{N['ho_all']} commands, the signalling log records"),

        ("Logistic regression. A regularised linear model over the 107 features, included as the simplest learner that can use the full feature set.",
         f"Logistic regression. A regularised linear model over the {N['n_features']} features, included as the simplest learner that can use the full feature set. It is worth noting ahead of Table 5.5 that this arm is slightly ahead of untuned gradient boosting under the fixed-hyper-parameter ladder. That is a property of the default settings rather than of the model families: a linear decision boundary has little capacity to fit the campaign it was trained on, so it loses least when the test campaign is a different corridor. Under the tuning budget of Table 5.4 the ordering reverses, which is itself the point of Section 5.3 -- the protocol reorders the leaderboard, so a leaderboard is not a property of the learners alone."),

        ("The branching ratio α/β is the expected number of direct offspring per event and "
         "summarises the strength of the clustering.",
         "The branching ratio α/β is the expected number of direct offspring per event under "
         "this kernel, and summarises the strength of the clustering. Section 5.10 reports that "
         "the Ogata residual test rejects the exponential kernel on this data, so that reading is "
         "withdrawn there and the ratio is used only as a comparable summary statistic."),
    ]

DELETE_EXACT = [
    "R_d(λ) = \t{ t ∈ d : y_t = 1 ∧ p̂_t < λ }",
    "and R̂(λ) averages over calibration drives rather than over samples.",
    ("Figure 4.1 illustrates the two formulations side by side.", 2),
]

def apply_fixups(doc, N):
    paras = list(doc.paragraphs)
    for spec in fixups(N):
        old, new, want = (*spec, 1) if len(spec) == 2 else spec
        hits = [p for p in paras if old in p.text]
        assert len(hits) == want, \
            f"fixup matched {len(hits)} times, expected {want}: {old[:70]!r}"
        for p in hits:
            set_text(p, p.text.replace(old, new))
    for spec in DELETE_EXACT:
        txt, found = (spec, 1) if isinstance(spec, str) else spec
        hits = [p for p in paras if p.text.strip() == txt]
        if found == 1:
            for p in hits:
                delete(p)
        elif len(hits) > 1:
            for p in hits[1:]:
                delete(p)

def main():
    print("Starting build...")
    N = build()
    doc = docx.Document(str(SRC))
    P = list(doc.paragraphs)
    LM = line_map(doc)
    raw = {**part_a(N), **part_b(N), **part_c(N)}
    ops = {(LM[k] if k >= 188 else k): v for k, v in raw.items()}
    assert len(ops) == len(raw), "paragraph key collision"
    TB = tables(N)

    # ---- tables
    print("Writing tables...")
    for slot, key in TABLE_SLOT.items():
        fill_table(doc.tables[slot], TB[key], doc)
        set_text(P[TABLE_CAPTION_PARA[slot]], TABLE_CAPTIONS[key])
    anchors = {key: doc.tables[slot]._tbl for key, slot in INSERT_AFTER_TABLE.items()}
    for key, slot in INSERT_AFTER_TABLE.items():
        anchor = anchors[key]
        cap_src = P[TABLE_CAPTION_PARA[slot]]
        cap_el = copy.deepcopy(cap_src._p)
        anchor.addnext(cap_el)
        cap = docx.text.paragraph.Paragraph(cap_el, cap_src._parent)
        set_text(cap, TABLE_CAPTIONS[key])
        new_table(doc, TB[key], cap)

    # ---- figures
    print("Writing figures...")
    for idx, (png, caption) in FIGURES.items():
        cap_par = next(P[j] for j in range(idx + 1, idx + 4) if P[j].text.strip())
        if png is DELETE:
            delete(P[idx])
            delete(cap_par)
            continue
        if png:
            p = FIGS / png
            if p.exists():
                replace_image(doc, P[idx], p)
            else:
                print("missing figure", png)
        if caption:
            set_text(cap_par, caption)

    # ---- text
    print("Writing text...")
    for idx, val in ops.items():
        p = P[idx]
        if val is DELETE:
            delete(p)
        elif isinstance(val, list):
            set_text(p, val[0])
            insert_after(p, val[1:])
        else:
            set_text(p, val)

    # ---- references
    print("Writing references...")
    ref_start, ref_end = LM[591], LM[624]
    slots = ref_end - ref_start + 1
    for i, (_, txt) in enumerate(REFS[:slots]):
        set_text(P[ref_start + i], f"[{i + 1}] {txt}")
    if len(REFS) > slots:
        insert_after(P[ref_end], [f"[{i + 1}] {txt}"
                                  for i, (_, txt) in enumerate(REFS) if i >= slots])
    else:
        for j in range(ref_start + len(REFS), ref_end + 1):
            delete(P[j])

    # ---- lists of figures and tables
    print("Writing lists of figures and tables...")
    def _num(cap):
        import re as _re
        m = _re.match(r"(?:Table|Figure) ([0-9A-D]+)\.([0-9]+)([a-z]?)", cap)
        ch, n, suf = m.group(1), int(m.group(2)), m.group(3)
        order = {"A": 100, "B": 101, "C": 102, "D": 103}.get(ch, int(ch) if ch.isdigit() else 999)
        return (order, n, suf)

    figs = sorted([c for _, c in (FIGURES[k] for k in sorted(FIGURES)) if c and c is not DELETE],
                  key=_num)
    if (FIGS / "fig_c20_cases.png").exists():
        figs.append(NEW_FIG_CAPTION)
        figs.sort(key=_num)
    lof = P[151:179]
    for i, p in enumerate(lof):
        if i < len(figs):
            set_text(p, figs[i])
        else:
            delete(p)
    if len(figs) > len(lof):
        insert_after(P[178], figs[len(lof):])

    LOT_RANGE = range(180, 206)
    ch6 = [P[i].text for i in LOT_RANGE if P[i].text.startswith("Table 6.")]
    tab_caps = sorted(list(TABLE_CAPTIONS.values()) + ch6, key=_num)
    lot = [P[i] for i in LOT_RANGE]
    for i, p in enumerate(lot):
        if i < len(tab_caps):
            set_text(p, tab_caps[i])
        else:
            delete(p)
    extra = tab_caps[len(lot):]
    if extra:
        insert_after(P[LOT_RANGE.stop - 1], extra)

    # ---- appendix D
    print("Writing Appendix D...")
    last = P[-1]
    hd = insert_after(last, ["APPENDIX D", "THE 22 AUDITED MODELS",
                             "Table D.1 lists every model audited in Section 2.3, so that each row of "
                             "Table 2.1 can be checked. The screening set was 108 publications; these 22 are "
                             "those presenting a learned or analytical model that predicts handover, radio-link "
                             "failure or next-cell occupancy.",
                             TABLE_CAPTIONS["D.1"]])
    hd[1].style = P[725].style
    hd[3].style = P[TABLE_CAPTION_PARA[26]].style
    new_table(doc, appendix_d(), hd[3])

    # ---- Figure 5.20
    print("Writing Figure 5.20...")
    cap_src = P[TABLE_CAPTION_PARA[18]]
    cands = [q for q in doc.paragraphs if q.text.startswith("Table 5.16 Nine cases")]
    anchor = cands[-1]
    assert anchor._p.getnext() is not None and anchor._p.getnext().tag.endswith("}tbl"), \
        "Figure 5.20 anchor is not the body caption of Table 5.16"
    fig_png = FIGS / "fig_c20_cases.png"
    if fig_png.exists():
        el = anchor._p.getnext()
        newp = copy.deepcopy(anchor._p)
        el.addnext(newp)
        par = docx.text.paragraph.Paragraph(newp, anchor._parent)
        for child in list(par._p):
            par._p.remove(child)
        run = par.add_run()
        w, h = Image.open(fig_png).size
        width = Inches(6.1)
        run.add_picture(str(fig_png), width=width, height=Emu(int(width * h / w)))
        par.alignment = 1
        cap_el = copy.deepcopy(cap_src._p)
        par._p.addnext(cap_el)
        set_text(docx.text.paragraph.Paragraph(cap_el, cap_src._parent), NEW_FIG_CAPTION)
    else:
        print("missing figure fig_c20_cases.png")

    print("Applying fixups...")
    apply_fixups(doc, N)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("Successfully built and saved to", OUT)

    MIRROR.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(OUT, MIRROR)
    print("Successfully mirrored to", MIRROR)

if __name__ == "__main__":
    main()
